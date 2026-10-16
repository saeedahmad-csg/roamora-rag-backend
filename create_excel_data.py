import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def create_ecommerce_products_excel(output_path="data/products.xlsx"):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Store_Product_Catalog"
    
    # Headers
    headers = [
        "Product SKU",
        "Product Name",
        "Category",
        "Regular Price ($)",
        "Sale Price ($)",
        "Stock Status & Inventory",
        "Key Specifications & Features",
        "Warranty & Return Policy",
        "Customer Rating & Reviews"
    ]
    
    # Product items for Lumina Direct E-Commerce Store
    products = [
        [
            "LUM-AUD-01",
            "LuminaSound Elite Wireless Headphones",
            "Audio & Sound",
            "$249.99",
            "$199.99",
            "In Stock (85 units available)",
            "Hybrid Active Noise Cancellation (up to 38dB), 40-hour battery life, USB-C fast charge (10 mins = 4 hours), multipoint Bluetooth 5.3, custom 40mm beryllium drivers, fold-flat design with carrying case",
            "1-Year Manufacturer Warranty included. 30-Day Hassle-Free Returns. Eligible for LuminaCare 2-Year Drop & Spill Protection ($29.99).",
            "4.8 / 5.0 (1,420 customer reviews)"
        ],
        [
            "LUM-WCH-02",
            "LuminaPulse Ultra Smartwatch",
            "Wearables & Fitness",
            "$179.99",
            "$149.99",
            "In Stock (140 units available)",
            "1.43-inch AMOLED Always-On Display, optical heart rate, ECG rhythm tracking, Blood Oxygen (SpO2), 110+ sports modes, 5ATM water resistant (swim-proof up to 50m), 7-day battery life, compatible with iOS & Android",
            "1-Year Manufacturer Warranty. 30-Day Money-Back Guarantee. LuminaCare 3-Year Battery & Accidental Plan ($49.99).",
            "4.7 / 5.0 (980 customer reviews)"
        ],
        [
            "LUM-SMH-03",
            "LuminaGlow Smart Ambient LED Light Strip",
            "Smart Home Lighting",
            "$49.99",
            "$39.99",
            "In Stock (310 units available)",
            "16.4 ft (5 meters) length, 16 million RGB colors + tunable warm-to-cool whites, Wi-Fi 2.4GHz & Bluetooth app control, works seamlessly with Amazon Alexa, Apple HomeKit & Google Assistant, built-in mic for music rhythmic synchronization",
            "1-Year Replacement Warranty. 30-Day Return window (must be unused in original packaging).",
            "4.6 / 5.0 (2,150 customer reviews)"
        ],
        [
            "LUM-ACC-04",
            "LuminaVault 65W GaN Fast Charger",
            "Charging & Accessories",
            "$39.99",
            "$29.99",
            "In Stock (420 units available)",
            "Next-gen Gallium Nitride (GaN) technology, 3 ports (2x USB-C Power Delivery 3.0 + 1x USB-A QuickCharge 4.0), foldable prongs, compact travel size (50% smaller than standard chargers), powers MacBook, iPhone, Galaxy, iPad simultaneously",
            "2-Year Extended Hardware Warranty included free. 30-Day Return Window.",
            "4.9 / 5.0 (3,410 customer reviews)"
        ],
        [
            "LUM-CAM-05",
            "LuminaVision 2K Outdoor Security Camera",
            "Smart Home Security",
            "$89.99",
            "$69.99",
            "In Stock (115 units available)",
            "2K QHD resolution (2560x1440), full color night vision with dual motion spotlights, 2-way audio intercom, AI person & vehicle detection, IP65 waterproof weatherproof rating, solar-panel compatible, microSD local storage (up to 256GB) or cloud",
            "1-Year Manufacturer Warranty. 30-Day Free Returns. Free replacement if damaged by weather within warranty.",
            "4.6 / 5.0 (740 customer reviews)"
        ],
        [
            "LUM-KEY-06",
            "LuminaTactile Pro Wireless Mechanical Keyboard",
            "Computer Accessories",
            "$129.99",
            "$109.99",
            "In Stock (65 units available)",
            "75% compact layout (84 keys), hot-swappable linear red switches, PBT double-shot keycaps, south-facing per-key RGB backlighting, tri-mode connectivity (Bluetooth 5.1, 2.4GHz wireless dongle, USB-C wired), Mac & Windows compatible",
            "1-Year Limited Warranty. 30-Day Return Policy. Switch puller and spare keycaps included in box.",
            "4.8 / 5.0 (610 customer reviews)"
        ],
        [
            "LUM-HUB-07",
            "LuminaDock 10-in-1 USB-C Hub & Docking Station",
            "Computer Accessories",
            "$79.99",
            "$64.99",
            "In Stock (200 units available)",
            "Dual 4K HDMI video output (dual monitor support 60Hz), 100W Power Delivery pass-through charging, Gigabit Ethernet (1000Mbps), SD/TF 4.0 card reader (up to 312MB/s), 3x USB 3.2 Gen 1 ports (5Gbps), precision aluminum unibody enclosure",
            "1-Year Hardware Warranty. 30-Day Return Policy. Protected against overheating and voltage spikes.",
            "4.7 / 5.0 (1,120 customer reviews)"
        ],
        [
            "LUM-ROB-08",
            "LuminaClean X7 Robot Vacuum & Sonic Mop",
            "Smart Appliances",
            "$399.99",
            "$329.99",
            "Low Stock (42 units available)",
            "Precision LiDAR laser mapping, 4000Pa ultra-strong suction, sonic vibration wet mopping (3000 scrubs/min), automatic ultrasonic carpet detection & mop auto-lift, 5200mAh battery (up to 180 min runtime), automatic self-recharging and resume",
            "1-Year Comprehensive Warranty (motors, sensors & battery). 30-Day Return Window. LuminaCare 3-Year Plan available ($49.99).",
            "4.8 / 5.0 (890 customer reviews)"
        ],
        [
            "LUM-SPK-09",
            "LuminaBoom 360 Portable Waterproof Speaker",
            "Audio & Sound",
            "$69.99",
            "$54.99",
            "In Stock (180 units available)",
            "360-degree omnidirectional sound, dual passive bass radiators, IPX7 100% waterproof (floats in pool/lake), 24-hour non-stop battery life, True Wireless Stereo (TWS) pairing allows syncing two units, shockproof rubberized bumper",
            "1-Year Limited Warranty. 30-Day Return Policy. Free return shipping for store credit.",
            "4.9 / 5.0 (2,850 customer reviews)"
        ],
        [
            "LUM-AIR-10",
            "LuminaPure H13 Smart HEPA Air Purifier",
            "Smart Appliances",
            "$119.99",
            "$99.99",
            "In Stock (90 units available)",
            "Medical-grade H13 True HEPA 4-stage filtration, removes 99.97% of airborne dust, pollen, pet dander, smoke and odor, covers up to 1,000 sq ft, whisper-quiet sleep mode (22dB), real-time laser PM2.5 air quality sensor with 4-color light ring",
            "2-Year Warranty on motor & electronics. 30-Day Money-Back Guarantee. Replacement filters available for $24.99.",
            "4.7 / 5.0 (1,340 customer reviews)"
        ]
    ]
    
    # Styling
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0D5C75", end_color="0D5C75", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    row_font = Font(name="Calibri", size=10)
    row_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    # Write header
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 28
    
    # Write data
    alt_fill = PatternFill(start_color="F0F8FA", end_color="F0F8FA", fill_type="solid")
    for r_idx, row_data in enumerate(products, start=2):
        ws.append(row_data)
        fill = alt_fill if r_idx % 2 == 0 else PatternFill(fill_type=None)
        for col_num in range(1, len(row_data) + 1):
            cell = ws.cell(row=r_idx, column=col_num)
            cell.font = row_font
            cell.alignment = row_align
            cell.border = thin_border
            if fill.fill_type:
                cell.fill = fill
        ws.row_dimensions[r_idx].height = 48

    # Column widths
    col_widths = [16, 32, 22, 18, 16, 26, 50, 38, 28]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
        
    wb.save(output_path)
    print(f"E-commerce product catalog Excel file successfully created at {output_path}")

if __name__ == "__main__":
    create_ecommerce_products_excel()
